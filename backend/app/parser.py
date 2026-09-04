import json
import os
import re
import shutil
import statistics
import subprocess
import tempfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pymupdf as fitz
from docx import Document
from docx.table import Table

from .config import Settings


class ParseError(ValueError):
    pass


@dataclass
class ParsedBlock:
    block_type: str
    content: str
    search_text: str
    heading_level: int | None = None
    page_number: int | None = None
    location_confidence: float = 0
    section_path: str | None = None


@dataclass
class ParsedDocument:
    page_count: int
    blocks: list[ParsedBlock]
    warnings: list[str] = field(default_factory=list)

    @property
    def text_char_count(self) -> int:
        return sum(len(block.search_text) for block in self.blocks)


def validate_file(data: bytes, suffix: str, settings: Settings) -> None:
    if not data:
        raise ParseError("文件为空")
    if len(data) > settings.max_upload_bytes:
        raise ParseError(f"文件超过 {settings.app_max_upload_mb} MB 限制")
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ParseError("文件内容不是有效 PDF")
        return
    if suffix != ".docx":
        raise ParseError("仅支持 DOCX 和文本型 PDF")
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ParseError("文件内容不是有效 DOCX")
            if len(names) > 5_000:
                raise ParseError("DOCX 内部文件数量超限")
            unpacked = sum(item.file_size for item in archive.infolist())
            if unpacked > 100 * 1024 * 1024:
                raise ParseError("DOCX 解压后体积超限")
            if any(".." in Path(name).parts or Path(name).is_absolute() for name in names):
                raise ParseError("DOCX 包含不安全路径")
    except BadZipFile as exc:
        raise ParseError("文件内容不是有效 DOCX") from exc


def find_libreoffice(settings: Settings) -> str | None:
    candidates = [
        settings.app_libreoffice_path,
        shutil.which("soffice") or "",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    return next((path for path in candidates if path and Path(path).is_file()), None)


def parse_document(path: Path, suffix: str, settings: Settings) -> ParsedDocument:
    result = parse_pdf(path, settings) if suffix == ".pdf" else parse_docx(path, settings)
    if result.text_char_count > settings.app_max_text_chars:
        raise ParseError("文档可提取文本超过处理上限")
    if not result.blocks or result.text_char_count < 20:
        raise ParseError("未提取到足够文本；扫描 PDF 暂不支持 OCR")
    return result


def parse_pdf(path: Path, settings: Settings) -> ParsedDocument:
    try:
        pdf = fitz.open(path)
    except Exception as exc:
        raise ParseError("PDF 无法打开或已损坏") from exc
    with pdf:
        if pdf.is_encrypted:
            raise ParseError("暂不支持加密 PDF")
        if not 0 < pdf.page_count <= settings.app_max_pages:
            raise ParseError(f"PDF 页数必须在 1 到 {settings.app_max_pages} 页之间")

        blocks: list[ParsedBlock] = []
        headings: list[str] = []
        warnings: list[str] = []
        for page_index, page in enumerate(pdf):
            page_items: list[tuple[float, float, ParsedBlock]] = []
            table_boxes: list[fitz.Rect] = []
            try:
                tables = page.find_tables().tables
            except Exception:
                tables = []
                warnings.append(f"第 {page_index + 1} 页表格识别失败，已保留普通文本")
            for table in tables:
                rows = table.extract()
                if not rows:
                    continue
                table_boxes.append(fitz.Rect(table.bbox))
                clean_rows = [[cell or "" for cell in row] for row in rows]
                search_text = "\n".join("\t".join(row) for row in clean_rows).strip()
                if search_text:
                    page_items.append(
                        (
                            table.bbox[1],
                            table.bbox[0],
                            ParsedBlock(
                                block_type="table",
                                content=json.dumps(clean_rows, ensure_ascii=False),
                                search_text=search_text,
                                page_number=page_index + 1,
                                location_confidence=1,
                            ),
                        )
                    )

            raw = page.get_text("dict", sort=True)
            text_blocks = [
                block
                for block in raw.get("blocks", [])
                if block.get("type") == 0
                and not any(
                    table_box.contains(fitz.Rect(block["bbox"]).tl)
                    and table_box.contains(fitz.Rect(block["bbox"]).br)
                    for table_box in table_boxes
                )
            ]
            sizes = [
                span["size"]
                for block in text_blocks
                for line in block.get("lines", [])
                for span in line.get("spans", [])
                if span.get("text", "").strip()
            ]
            body_size = statistics.median(sizes) if sizes else 10
            for raw_block in text_blocks:
                box = fitz.Rect(raw_block["bbox"])
                spans = [
                    span
                    for line in raw_block.get("lines", [])
                    for span in line.get("spans", [])
                    if span.get("text", "").strip()
                ]
                text = "\n".join(
                    "".join(span.get("text", "") for span in line.get("spans", [])).strip()
                    for line in raw_block.get("lines", [])
                ).strip()
                if not text:
                    continue
                max_size = max((span["size"] for span in spans), default=body_size)
                bold = any(span.get("flags", 0) & 16 for span in spans)
                numbered_heading = bool(
                    re.match(
                        r"^(?:第[一二三四五六七八九十\d]+[章节部分]|[一二三四五六七八九十]+[、.]|\d+(?:\.\d+)*[、.\s])",
                        text,
                    )
                )
                is_heading = len(text) <= 80 and (
                    max_size >= body_size * 1.22 or (bold and numbered_heading)
                )
                level = 1 if max_size >= body_size * 1.55 else 2 if is_heading else None
                block_type = "heading" if is_heading else "paragraph"
                page_items.append(
                    (
                        box.y0,
                        box.x0,
                        ParsedBlock(
                            block_type=block_type,
                            content=text,
                            search_text=text,
                            heading_level=level,
                            page_number=page_index + 1,
                            location_confidence=1,
                        ),
                    )
                )

            for _, _, block in sorted(page_items, key=lambda item: (item[0], item[1])):
                if block.block_type == "heading":
                    level = block.heading_level or 2
                    headings[level - 1 :] = [block.search_text]
                block.section_path = "/".join(headings) or None
                blocks.append(block)
        return ParsedDocument(page_count=pdf.page_count, blocks=blocks, warnings=warnings)


def parse_docx(path: Path, settings: Settings) -> ParsedDocument:
    soffice = find_libreoffice(settings)
    if not soffice:
        raise ParseError("解析 DOCX 页码需要安装 LibreOffice，或配置 APP_LIBREOFFICE_PATH")
    try:
        document = Document(path)
    except Exception as exc:
        raise ParseError("DOCX 无法打开或已损坏") from exc

    blocks: list[ParsedBlock] = []
    headings: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            search_text = "\n".join("\t".join(row) for row in rows).strip()
            if search_text:
                blocks.append(
                    ParsedBlock(
                        block_type="table",
                        content=json.dumps(rows, ensure_ascii=False),
                        search_text=search_text,
                        section_path="/".join(headings) or None,
                    )
                )
            continue

        text = item.text.strip()
        if not text:
            continue
        style_name = item.style.name if item.style else ""
        match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
        level = min(int(match.group(1)), 9) if match else None
        if level:
            headings[level - 1 :] = [text]
        blocks.append(
            ParsedBlock(
                block_type="heading" if level else "paragraph",
                content=text,
                search_text=text,
                heading_level=level,
                section_path="/".join(headings) or None,
            )
        )

    page_texts = _convert_docx_pages(path, soffice)
    _assign_docx_pages(blocks, page_texts)
    missing = sum(block.page_number is None for block in blocks)
    warnings = []
    if missing:
        warnings.append(f"{missing} 个结构块未能可靠匹配页码，相关块不可作为最终证据")
    return ParsedDocument(page_count=len(page_texts), blocks=blocks, warnings=warnings)


def _convert_docx_pages(path: Path, soffice: str) -> list[str]:
    with tempfile.TemporaryDirectory(prefix="bidguard-docx-") as temp:
        temp_dir = Path(temp)
        profile_dir = temp_dir / "profile"
        command = [
            soffice,
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(temp_dir),
            str(path),
        ]
        creation_flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=60,
                creationflags=creation_flags,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise ParseError("DOCX 转 PDF 超时") from exc
        pdf_path = temp_dir / f"{path.stem}.pdf"
        if result.returncode or not pdf_path.exists():
            raise ParseError("LibreOffice 无法将 DOCX 转换为 PDF")
        with fitz.open(pdf_path) as pdf:
            return [page.get_text("text", sort=True) for page in pdf]


def _normalize_location_text(value: str) -> str:
    return re.sub(r"[\W_]+", "", value, flags=re.UNICODE).lower()


def _assign_docx_pages(blocks: list[ParsedBlock], page_texts: list[str]) -> None:
    normalized_pages = [_normalize_location_text(text) for text in page_texts]
    cursor = 0
    for block in blocks:
        text = _normalize_location_text(block.search_text)
        if len(text) < 4:
            continue
        exact_pages = [
            index for index, page in enumerate(normalized_pages) if text in page
        ]
        needle = text[: min(60, len(text))]
        prefix_pages = [
            index for index, page in enumerate(normalized_pages) if needle in page
        ]
        candidates = exact_pages or prefix_pages
        if not candidates:
            continue
        page_index = next((index for index in candidates if index >= cursor), candidates[0])
        block.page_number = page_index + 1
        block.location_confidence = 1 if exact_pages else 0.85
        cursor = page_index
