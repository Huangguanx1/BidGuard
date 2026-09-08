"""Export only immutable, reviewed snapshots. No model calls on download."""
import json
from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

LABELS = {'high': '高', 'medium': '中', 'low': '低', 'pending': '待确认', 'confirmed': '已确认',
          'ignored': '已忽略', 'modified': '已修改', 'satisfied': '满足', 'partial': '部分满足',
          'not_satisfied': '不满足', 'not_found': '未找到', 'uncertain': '待确认'}
DISCLAIMER = '本报告仅供辅助审阅，不构成法律、招标或投标专业意见。模型可能遗漏或误判，请结合原始文件和专业判断使用。'


def build_snapshot(db, review_id):
    row = db.get_review(review_id)
    findings = [f.model_dump() for f in db.list_findings(review_id)]
    for finding in findings:
        finding['effective'] = {key: (finding['override'] or {}).get(key, finding[key])
                                for key in ('title', 'description', 'risk_level', 'suggestion')}
    return dict(schema_version='1.0', review_id=review_id, name=row['name'], completed_at=datetime.now(UTC).isoformat(),
        created_at=row['created_at'], model_name=row['model_name'], model_provider=row['model_provider'],
        prompt_version=row['prompt_version'], prompt_hash=row['prompt_hash'], model_parameters=json.loads(row['model_parameters']),
        input_tokens=row['input_tokens'], output_tokens=row['output_tokens'],
        documents={role: db.get_document(row[f'{role}_document_id']).model_dump() for role in ('tender', 'bid')},
        requirements=[r.model_dump() for r in db.list_requirements(review_id)],
        requirement_checks=[c.model_dump() for c in db.list_requirement_checks(review_id)], findings=findings,
        summary={'active_findings': sum(f['review_status'] != 'ignored' for f in findings),
                 'ignored_findings': sum(f['review_status'] == 'ignored' for f in findings),
                 'high_risk_findings': sum(f['review_status'] != 'ignored' and f['effective']['risk_level'] == 'high' for f in findings)},
        disclaimer=DISCLAIMER)


def report_bytes(snapshot, format):
    if format == 'json':
        return json.dumps(snapshot, ensure_ascii=False, indent=2).encode('utf-8'), 'application/json'
    doc = Document()
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.add_paragraph('招投标文件审查报告', 'Title')
    doc.add_paragraph(snapshot['name'])
    summary = snapshot['summary']
    doc.add_paragraph(f"本次审查共核对 {len(snapshot['requirements'])} 条要求。人工复核后保留 {summary['active_findings']} 项问题，其中高风险 {summary['high_risk_findings']} 项，另有 {summary['ignored_findings']} 项已忽略。以下记录包含原文证据和人工处置结果。")
    completed_label = datetime.fromisoformat(snapshot['completed_at']).strftime('%Y-%m-%d %H:%M UTC')
    doc.add_paragraph(f"招标文件：{snapshot['documents']['tender']['original_name']}\n投标文件：{snapshot['documents']['bid']['original_name']}\n完成时间：{completed_label}\n审查模型：{snapshot['model_name']}")
    doc.add_heading('人工复核结果', 1)
    if not snapshot['findings']:
        doc.add_paragraph('未发现需要处理的问题。这不代表文件不存在其他风险。')
    for index, finding in enumerate(snapshot['findings'], 1):
        final = finding['effective']
        doc.add_heading(f"{index} {final['title']}", 2)
        doc.add_paragraph(f"风险等级：{LABELS[final['risk_level']]}    人工状态：{LABELS[finding['review_status']]}")
        doc.add_paragraph(final['description'])
        doc.add_paragraph('建议：' + final['suggestion'])
        if finding['reviewer_note']:
            doc.add_paragraph('人工备注：' + finding['reviewer_note'])
        if finding['override']:
            doc.add_paragraph('原始结论：' + finding['title'] + '；' + finding['description'])
        for item in finding['evidence']:
            role = '招标' if item['document_role'] == 'tender' else '投标'
            warning = '（页码定位置信度较低）' if item.get('location_confidence', 1) < .9 else ''
            doc.add_paragraph(f"{role}文件 第 {item['page_number']} 页{warning}\n{readable_excerpt(item['excerpt'])}\n证据块 {item['block_id']}")
    doc.add_heading('逐项要求核对', 1)
    checks = {c['requirement_id']: c for c in snapshot['requirement_checks']}
    for index, requirement in enumerate(snapshot['requirements'], 1):
        check = checks.get(requirement['id'])
        doc.add_heading(f"{index} {requirement['title']}", 2)
        doc.add_paragraph(requirement['description'])
        if check:
            doc.add_paragraph(f"匹配结果：{LABELS[check['match_status']]}\n判断理由：{check['reason']}")
            for item in check['tender_evidence'] + check['bid_evidence']:
                doc.add_paragraph(f"{'招标' if item['document_role'] == 'tender' else '投标'} 第 {item['page_number']} 页：{readable_excerpt(item['excerpt'])}")
            if check['match_status'] == 'not_found':
                doc.add_paragraph('已检索投标块：' + '、'.join(check['searched_block_ids']))
    doc.add_heading('使用说明', 1)
    doc.add_paragraph(snapshot['disclaimer'])
    doc.add_paragraph(f"审查编号：{snapshot['review_id']}\n提示词版本：{snapshot['prompt_version']}\nToken 用量：{snapshot['input_tokens'] + snapshot['output_tokens']}")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('BidGuard AI   ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    timestamp = datetime.fromisoformat(snapshot['completed_at'])
    doc.core_properties.created = doc.core_properties.modified = timestamp
    out = BytesIO()
    doc.save(out)
    # Fix ZIP metadata so identical frozen snapshots produce identical downloads.
    normalized = BytesIO()
    with ZipFile(BytesIO(out.getvalue())) as source, ZipFile(normalized, 'w', ZIP_DEFLATED) as destination:
        for entry in source.infolist():
            entry.date_time = (2020, 1, 1, 0, 0, 0)
            destination.writestr(entry, source.read(entry.filename))
    return normalized.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def readable_excerpt(text):
    try:
        value = json.loads(text)
        if isinstance(value, list) and all(isinstance(row, list) for row in value):
            return '\n'.join(' | '.join(str(cell) for cell in row) for row in value)
    except (ValueError, TypeError):
        pass
    return text
