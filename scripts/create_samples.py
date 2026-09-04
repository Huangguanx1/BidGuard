from pathlib import Path

import pymupdf as fitz
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


def find_chinese_font() -> Path:
    candidates = [
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
    ]
    font = next((path for path in candidates if path.is_file()), None)
    if not font:
        raise RuntimeError("生成中文 PDF 样本需要微软雅黑、Noto Sans CJK 或苹方字体")
    return font


def create_docx() -> None:
    document = Document()
    document.add_heading("星河市数字档案平台招标文件", level=1)
    document.add_paragraph("本文件中的单位、项目和数据均为虚构内容，仅用于软件演示。")
    document.add_heading("一、资格审查", level=2)
    document.add_paragraph("投标人须提供有效的营业执照复印件并加盖公章。")
    document.add_heading("二、工期要求", level=2)
    document.add_paragraph("项目建设工期不得超过150日历天。")
    document.add_heading("三、评分标准", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "评分项"
    table.rows[0].cells[1].text = "分值"
    table.rows[0].cells[2].text = "要求"
    row = table.add_row().cells
    row[0].text = "实施方案"
    row[1].text = "20"
    row[2].text = "内容完整、计划合理"
    document.save(SAMPLES / "fictional_tender.docx")


def create_pdf(
    path: Path,
    title: str,
    lines: list[str],
    table_rows: list[list[str]] | None = None,
) -> None:
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    font_name = "sample-cjk"
    page.insert_font(fontname=font_name, fontfile=str(find_chinese_font()))
    page.insert_text((64, 72), title, fontsize=18, fontname=font_name)
    y = 116
    for line in lines:
        page.insert_textbox(
            fitz.Rect(64, y, 530, y + 54),
            line,
            fontsize=11,
            fontname=font_name,
        )
        y += 62
    if table_rows:
        widths = [210, 70, 190]
        row_height = 30
        x_positions = [64]
        for width in widths:
            x_positions.append(x_positions[-1] + width)
        for row_index in range(len(table_rows) + 1):
            line_y = y + row_index * row_height
            page.draw_line((x_positions[0], line_y), (x_positions[-1], line_y))
        for x in x_positions:
            page.draw_line((x, y), (x, y + len(table_rows) * row_height))
        for row_index, row in enumerate(table_rows):
            for column_index, cell in enumerate(row):
                page.insert_textbox(
                    fitz.Rect(
                        x_positions[column_index] + 4,
                        y + row_index * row_height + 6,
                        x_positions[column_index + 1] - 4,
                        y + (row_index + 1) * row_height - 2,
                    ),
                    cell,
                    fontsize=9,
                    fontname=font_name,
                )
    pdf.subset_fonts()
    pdf.save(path)


def main() -> None:
    SAMPLES.mkdir(exist_ok=True)
    create_docx()
    create_pdf(
        SAMPLES / "fictional_tender.pdf",
        "星河市数字档案平台招标文件",
        [
            "资格审查：投标人须提供有效的营业执照复印件并加盖公章。",
            "工期要求：项目建设工期不得超过150日历天。",
            "评分标准：实施方案内容完整、计划合理，满分20分。",
        ],
        [["评分项", "分值", "要求"], ["实施方案", "20", "内容完整、计划合理"]],
    )
    create_pdf(
        SAMPLES / "fictional_bid.pdf",
        "星云科技数字档案平台投标文件",
        [
            "本公司已提供有效营业执照复印件并加盖公章。",
            "项目名称：星河市数字档案平台。",
            "承诺建设工期为180日历天。",
            "投标报价为人民币980000元。",
        ],
    )
    print(f"Created fictional samples in {SAMPLES}")


if __name__ == "__main__":
    main()
